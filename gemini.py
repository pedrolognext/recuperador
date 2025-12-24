import io
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_TABLE_ALIGNMENT
# Cargar variables de entorno desde ubicación estándar
import env_loader

# --------------------------------------------------------------------------
# DICCIONARIOS DE TEXTOS POR IDIOMA
# --------------------------------------------------------------------------

TEXTS = {
    "es": {
        "professional_summary": "Resumen Profesional",
        "professional_experience": "Experiencia Profesional",
        "education": "Educación",
        "certifications": "Certificaciones",
        "languages": "Idiomas",
        "skills": "Skills",
        "reference": "Reference:",
        "page": "Página"
    },
    "en": {
        "professional_summary": "Professional Summary",
        "professional_experience": "Professional Experience",
        "education": "Education",
        "certifications": "Certifications",
        "languages": "Languages",
        "skills": "Skills",
        "reference": "Reference:",
        "page": "Page"
    }
}

# --------------------------------------------------------------------------
# CONFIGURACIÓN Y DATOS EXTRAÍDOS
# --------------------------------------------------------------------------

PROFILE_DATA = {
    "name": "LUIS JAVIER JIMÉNEZ NAVARRO",
    "role": "Senior Java Developer",
    "reference": "CAND7082 L.J.J.N.",
    "summary": (
        "Senior Java Developer located in Madrid with extensive experience in Java, Spring, Junit, TDD, "
        "C/C++, SQL, BackEnd, Scrum, Jira, GIT, Docker, Intellij, Sonar, Teamwork, and Finance. "
        "Cloud Experience includes AWS, Kubernetes, Firestore, GitLab, Github, gRPC, and Hexagonal Architecture."
    ),
    "languages": [
        {"lang": "Spanish", "level": "Native"},
        {"lang": "English", "level": "Working Proficiency"}
    ],
    "skills": [
        "Core: Java (6/7/8/17-21), C/C++, SQL, BackEnd, Hexagonal Architecture, Microservices.",
        "Frameworks & Libraries: Spring Boot, Junit, TDD, gRPC, Traffic Parrot, Quartz, Akka Actors.",
        "Cloud & DevOps: AWS, Kubernetes, OpenShift, Docker, GitLab, Github, Kafka, Redis, Firestore.",
        "Tools: Intellij, Sonar, Jira, GIT, Splunk, New Relic, Jenkins, Maven, Gradle.",
        "Databases: MongoDB, PostgreSQL, Oracle, SQLite, SQL Server 2008 R2.",
        "Methodologies: Scrum, Agile Squad, Domain Driven Design (DDD)."
    ]
}

EXPERIENCE_LIST = [
    {
        "company": "K-LAGAN",
        "role": "Senior Java Developer",
        "dates": "Aug 2024 - Current",
        "description": "Project: PSLAR (Inditex). Managed the consumption of images and their attributes for the Inditex store network, as well as updating information about the stores on Google. Stack: Java 17-21, Spring Boot, MongoDB, gRPC, OpenShift, Docker, GitHub, Kafka, Traffic Parrot, Redis.",
        "achievements": [
            "Significantly boosted project development and maintenance by integrating solutions for efficient communication between microservices, using modern technologies such as gRPC, or specific Inditex solutions based on Java."
        ]
    },
    {
        "company": "Innovbo Group",
        "role": "Expert Senior Developer",
        "dates": "Mar 2024 - Jul 2024",
        "description": "Expert Senior Developer role focused on code quality, mentoring, and third-party integrations within an Agile Squad environment.",
        "achievements": [
            "Performed code reviews, mentored and acted as part of the Agile Squad.",
            "Developed and tested a third-party integration with MonitorHub to process GPS geolocations."
        ]
    },
    {
        "company": "FSB Technology",
        "role": "Senior Java Developer",
        "dates": "Jun 2022 - Jan 2024",
        "description": "Project: Fitzdares Canada (June 2022-June 2023). Launch of the Fitzdares Canada website with online casino games and sports betting. Project: Support (June 2023-January 2024). Support for backend systems in production.",
        "achievements": [
            "Managed changes to Registration Form & Account Details on the Backend side with PostgreSQL.",
            "Tested and troubleshooted casino games from the providers Scientific Games and Evolution Games.",
            "Configured GeoComply and IDComply to check geolocation and authorization for compliance.",
            "Production release of the website fitzdares.ca and on-call activities as part of the Bootcamp Team.",
            "Troubleshooting issues found by the Datadog platform. Checked logs in the different production environments.",
            "Scaled any issue to the proper team once the source of the problem has been identified.",
            "Executed queries against the Postgres databases to diagnose issues."
        ]
    },
    {
        "company": "EPAM",
        "role": "Senior Software Engineer",
        "dates": "Apr 2021 - Jun 2022",
        "description": "Developed requirements for the project PCWS 'Price Change Web Service' from Nordstrom. Worked with Kubernetes, GitLab CI/CD, Spring Boot, Kafka, and various monitoring tools in an On-Call environment.",
        "achievements": [
            "Adapted Kubernetes configuration to use a Gitlab CI/CD standard pipeline.",
            "Changed configuration of a Spring Boot application to use yaml files instead of property files.",
            "Checked logs in CI/CD standard pipeline from GitLab and migrated Docker images from Gitlab repository to Artifactory.",
            "Changed libsonnet files to yml files for Kubernetes configuration and scaled in/out Kubernetes resources.",
            "Coded PoC microservices for reading from and writing to Kafka topics.",
            "Coded a multithreaded audit program for validating data coming from several REST service endpoints, troubleshooting deadlock issues, and handling concurrency.",
            "Checked logs in Splunk and analyzed charts in New Relic as part of the On-Call tasks.",
            "Refactored tasks coded in Python for the Airflow system and checked logs as part of the On-Call responsibility."
        ]
    },
    {
        "company": "SMARTMATIC",
        "role": "Java Senior Developer",
        "dates": "Sep 2019 - Sep 2020",
        "description": "Developed requirements on electoral systems for the countries: United States and Uganda. Programmed in Java 6, 7 or 8 with Spring Boot according to the assigned module, connecting to Oracle and SQLite databases. Managed Dockerized environments: Database, Java BackEnd, and AngularJS FrontEnd.",
        "achievements": [
            "Created a Microservice using Spring Boot for the Platform project within the framework of hexagonal architecture.",
            "Transformed a monolithic program into microservices by using the Strangler Pattern and Domain Driven Design.",
            "Used Intellij for debugging directly inside Docker containers, JRebel/Maven Helper for rapid deployments, and SonarLint for quality checks.",
            "Coded change controls for the Oracle database through XML files and handled DB Browser for SQLite.",
            "Version control with GIT (Merge Request/Code Review). Used Docker in Windows with Linux instances.",
            "Participated in Scrum ceremonies (Daily, Planning, Review, Grooming, Retrospective)."
        ]
    },
    {
        "company": "GLOBANT",
        "role": "Java Developer",
        "dates": "Nov 2017 - Aug 2019",
        "description": "Developed user stories for Southwest Airlines and Royal Caribbean client accounts, utilizing Java 8 with Spring Boot, Oracle database, Angular7 and React/Redux. Member of a Scrum team with full CI/CD pipeline management.",
        "achievements": [
            "Handled tools for continuous integration: Jenkins builds and Rundeck deployments.",
            "Checked quality and technical debt using a Sonar server.",
            "Created and refined unit tests with JUnit/Mockito (TDD) and behavioral tests with JBehave/Cucumber (BDD).",
            "Handled Atlassian suite: Jira, Confluence, and Stash (GIT repository).",
            "Developed user stories related to the Royal Caribbean client account: Apigee rules, Kafka (Akka Actors), and Kubernetes pod monitoring (kubectl).",
            "Created a proof of concept on minikube for a legacy Java application migration."
        ]
    },
    {
        "company": "TATA CONSULTANCY SERVICES",
        "role": "Senior Developer",
        "dates": "Mar 2016 - Oct 2017",
        "description": "Managed technologies related to the Banistmo project. Developed applications in Java 6/8 with IBM Websphere Application Server 8.5, Microsoft SQL Server 2008 R2, Angular2, and Apache Tomcat Server deployment.",
        "achievements": [
            "Developed queries and stored procedures in Microsoft SQL Server 2008 R2.",
            "Developed requirements in Java 6, deployed applications in IBM Websphere Application Server 8.5 using Harvest.",
            "Developed an application formed by several CRUDs built of RESTful services in Java 8, Front-End in Angular2, Maven dependencies, and Apache Tomcat Server deployment."
        ]
    },
    {
        "company": "SOFTWARE ESTRATÉGICO",
        "role": "Outsourced Developer",
        "dates": "Sep 2015 - Feb 2016",
        "achievements": [
            "Programmed in Java 6 for the execution of tasks using the Quartz library.",
            "Made Nginx setup as a reverse proxy for Apache for projects PILA and Pago de Cesantias (Enlace Operativo y Compuredes).",
            "Prepared reports using JavaServer Faces and Spring Web Flow (Company: Quipux)."
        ]
    },
    {
        "company": "CASTILLOMAX OIL AND GAS",
        "role": "Lead Programmer",
        "dates": "Mar 2014 - Jan 2015",
        "description": "Coordinated programming tasks of the OPTIMAX 3D Project, a real-time 3D simulator of marine terminals. Built with Unity 3D, C#, and used devices like Leap Motion Controller and Virtual Reality Glasses.",
        "achievements": [
            "Managed a multidisciplinary team of 7 people including engineers and programmers.",
            "Software presented at World Petroleum Congress (Russia) and Latin American Petroleum Show (LAPS 2014).",
            "Successfully delivered a cutting-edge 3D simulation system using advanced VR technologies."
        ]
    },
    {
        "company": "ECRS CORPORATION",
        "role": "Project Engineer",
        "dates": "Feb 2012 - Sep 2013",
        "achievements": [
            "Implemented Driver on C++ for the fiscal printer QPrint MF.",
            "Created demo programs on C++, VB6, Java and C# (.Net) for integration tasks.",
            "Developed system manuals and user manuals of the Driver QPrint MF.",
            "Managed backup tasks with Norton Ghost 15 and the Company's Website (PHP5/MySQL)."
        ]
    },
    {
        "company": "CLEVER FINANCIAL",
        "role": "Investment's System Manager",
        "dates": "Jan 2007 - Aug 2011",
        "achievements": [
            "Analysis, implementation and control of investment systems for financial assets.",
            "Updated portfolios current value through Bloomberg API, managed reporting and backup automatization.",
            "Tools: SQL Server 2008, VS.NET C# 2008, Matlab 7, Bloomberg API, iTextSharp, AJAX.",
            "Managed a team of 2 IT people."
        ]
    },
    {
        "company": "PRICEWATERHOUSE COOPER",
        "role": "Software Engineer",
        "dates": "Oct 2004 - Feb 2006",
        "achievements": [
            "Developed a Workflow System managing document approvals, parametrization of events, routes, rules, roles, and politics.",
            "Tools: SQL Server 2000, VS.NET C# 2003, Web Services, Windows Services, XML, MSMQ.",
            "Provided support on VB6 programs, Excel macros, SAP Gui, and Oracle Financial front end."
        ]
    },
    {
        "company": "INDEPENDENT CONSULTANT",
        "role": "Consultant",
        "dates": "Mar 2003 - Jun 2004",
        "achievements": [
            "IESA: Developed Web forms in ASP.NET with C# and Oracle 9i for Managers Evaluation Program.",
            "Ministry of Health: Developed Web forms in JSP and PostgreSQL 7.2 for the Oncological Program.",
            "Ingeniería Caura: Designed Access 2000 database and VBA reporting interfaces.",
            "Archicentro: Configured database for ERP program SCAV 2.5."
        ]
    },
    {
        "company": "PDVSA INTEVEP",
        "role": "Programmer Analyst",
        "dates": "Feb 2002 - Dec 2002",
        "achievements": [
            "Designed object oriented graphic libraries 2D and 3D in OpenGL and C++ language.",
            "Designed and developed VRJuggler applications for visualization of reservoirs on immersive displays (CAVE, Walls).",
            "Integrated modules in C, C++, Java and Fortran using JNI.",
            "Wrote technical documentation under UML notation (RUP Methodology).",
            "Migrated code for Windows 2000, IRIX and Solaris."
        ]
    },
    {
        "company": "PDVSA INTEVEP",
        "role": "Intern",
        "dates": "Jul 2001 - Dec 2001",
        "achievements": [
            "Designed and developed a system for oil fields visualization using OpenGL, Java, C++ for Windows, Solaris and Irix operating systems."
        ]
    },
    {
        "company": "IESA",
        "role": "Intern",
        "dates": "Jul 1999 - Sep 1999",
        "achievements": [
            "Upgraded the handling of activities on a control system for the Institute's Help Desk.",
            "Developed SQL queries along with an interface on the intranet system using ASP and JDBC connection."
        ]
    }
]

EDUCATION_LIST = [
    "Graduate Diploma in Education Component - ALEJANDRO DE HUMBOLDT UNIVERSITY (Aug 2008)",
    "Master in Finance - INSTITUTE OF SUPERIOR STUDIES IN ADMINISTRATION (IESA) (Nov 2004)",
    "Software Engineer - SIMÓN BOLIVAR UNIVERSITY (USB) (Mar 2002)"
]

CERTIFICATIONS_LIST = [
    "Udemy: Learn Docker from Scratch to Swarm and Kubernetes",
    "Udemy: Introducing Spring Boot",
    "Udemy: AngularJS For Beginners",
    "Udemy: Spring Boot. Complete guide from development to deployment",
    "Udemy: MongoDB for Java Developers - Project Based",
    "Udemy: Quick introduction to influxdb",
    "VRJuggler Programming (2002)",
    "Project Monitoring and Control using Earned Value (2014)",
    "Money Laundering Prevention (2009)"
]

# --------------------------------------------------------------------------
# FUNCIONES DE ESTILO Y GENERACIÓN
# --------------------------------------------------------------------------

def set_font_style(run, size=11, bold=False, color_rgb=None, underline=False):
    run.font.name = 'Space Grotesk'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.underline = underline
    if color_rgb:
        run.font.color.rgb = color_rgb
    else:
        run.font.color.rgb = RGBColor(0, 0, 10) # #00000A

def set_cell_background(cell, color_hex):
    """Aplica color de fondo a una celda de tabla"""
    from docx.oxml.shared import qn
    cell_xml_element = cell._tc
    table_cell_properties = cell_xml_element.get_or_add_tcPr()
    shade_obj = OxmlElement('w:shd')
    shade_obj.set(qn('w:fill'), color_hex)
    table_cell_properties.append(shade_obj)

def set_table_borders(table):
    """Aplica bordes negros a toda la tabla"""
    from docx.oxml.shared import qn
    
    # Configurar bordes de la tabla
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    
    # Crear elemento de bordes
    tbl_borders = OxmlElement('w:tblBorders')
    
    # Definir cada tipo de borde
    border_types = ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']
    
    for border_type in border_types:
        border = OxmlElement(f'w:{border_type}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')  # Grosor del borde
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')  # Color negro
        tbl_borders.append(border)
    
    tbl_pr.append(tbl_borders)

def setup_document_styles(doc):
    # Base Normal Style
    style = doc.styles['Normal']
    style.font.name = 'Space Grotesk'
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0, 0, 10)
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def add_header(doc, logo_path):
    section = doc.sections[0]
    section.header_distance = Cm(1.25)
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
   
    # Add Logo
    run = paragraph.add_run()
    run.add_picture(logo_path, height=Cm(1.3))

def add_page_number_field(paragraph):
    """Agrega un campo de número de página que funciona correctamente"""
    from docx.oxml.shared import qn
    
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')

    run_element = paragraph._p.add_r()
    run_element.append(fldChar1)
    run_element.append(instrText)
    run_element.append(fldChar2)
    
    return run_element

def add_total_pages_field(paragraph):
    """Agrega un campo de número total de páginas que funciona correctamente"""
    from docx.oxml.shared import qn
    
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'NUMPAGES'

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')

    run_element = paragraph._p.add_r()
    run_element.append(fldChar1)
    run_element.append(instrText)
    run_element.append(fldChar2)
    
    return run_element

def add_footer(doc, ref_code, language="es"):
    section = doc.sections[0]
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
   
    # Limpiar el párrafo existente
    paragraph.clear()
   
    # Agregar el texto de referencia
    page_text = TEXTS[language]["page"]
    run = paragraph.add_run(f"{ref_code} – {page_text} ")
    set_font_style(run, size=11, bold=True)
   
    # Agregar campo de página actual
    add_page_number_field(paragraph)
   
    # Agregar " de " / " of "
    separator = " de " if language == "es" else " of "
    run_sep = paragraph.add_run(separator)
    set_font_style(run_sep, size=11, bold=True)
   
    # Agregar campo de total de páginas
    add_total_pages_field(paragraph)

def create_experience_table(doc, exp_data):
    table = doc.add_table(rows=5, cols=1)  # Aumentamos a 5 filas para separar descripción y achievements
    table.autofit = False
    table.allow_autofit = False
   
    # Fila 1: Empresa
    cell_company = table.cell(0, 0)
    p = cell_company.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(exp_data['company'])
    set_font_style(run, size=11, bold=True)
   
    # Fila 2: Puesto
    cell_role = table.cell(1, 0)
    p = cell_role.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(exp_data['role'])
    set_font_style(run, size=11, bold=True, color_rgb=RGBColor(250, 60, 15)) # #FA3C0F
   
    # Fila 3: Fechas
    cell_date = table.cell(2, 0)
    p = cell_date.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(exp_data['dates'])
    set_font_style(run, size=11, color_rgb=RGBColor(102, 102, 102)) # #666666
   
    # Fila 4: Descripción (párrafo justificado)
    cell_description = table.cell(3, 0)
    p = cell_description.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.0
    
    # Agregar descripción si existe
    if 'description' in exp_data and exp_data['description']:
        run = p.add_run(exp_data['description'])
        set_font_style(run, size=11)
    else:
        p.text = "" # Si no hay descripción, dejar vacío
   
    # Fila 5: Achievements (viñetas)
    cell_achievements = table.cell(4, 0)
    p = cell_achievements.paragraphs[0]
    p.text = "" # Clean start
   
    # Agregar achievements si existen
    if 'achievements' in exp_data and exp_data['achievements']:
        for item in exp_data['achievements']:
            p = cell_achievements.add_paragraph(style='List Paragraph')
            p.paragraph_format.left_indent = Cm(0.63) # Standard bullet indentation
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(f"• {item}")
            set_font_style(run, size=11)
    elif 'desc' in exp_data:  # Fallback para experiencias que aún usan el formato antiguo
        for item in exp_data['desc']:
            p = cell_achievements.add_paragraph(style='List Paragraph')
            p.paragraph_format.left_indent = Cm(0.63)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(f"• {item}")
            set_font_style(run, size=11)
       
    # Remove empty first paragraph if it exists and is empty
    if len(cell_achievements.paragraphs) > 1 and not cell_achievements.paragraphs[0].text:
        p_element = cell_achievements.paragraphs[0]._element
        p_element.getparent().remove(p_element)

def generate_cv(logo_path, profile_data=None, experience_list=None, education_list=None, certifications_list=None, language="es", custom_reference=None):
    """
    Genera un CV en formato Word
    
    Args:
        logo_path: Ruta al archivo de logo
        profile_data: Diccionario con datos de perfil (opcional, usa PROFILE_DATA por defecto)
        experience_list: Lista de experiencias (opcional, usa EXPERIENCE_LIST por defecto)
        education_list: Lista de educación (opcional, usa EDUCATION_LIST por defecto)
        certifications_list: Lista de certificaciones (opcional, usa CERTIFICATIONS_LIST por defecto)
        language: Idioma del documento ("es" o "en")
        custom_reference: Referencia personalizada (sobrescribe la del profile_data)
    
    Returns:
        Document: Objeto documento de python-docx
    """
    # Usar valores por defecto si no se proporcionan
    if profile_data is None:
        profile_data = PROFILE_DATA
    if experience_list is None:
        experience_list = EXPERIENCE_LIST
    if education_list is None:
        education_list = EDUCATION_LIST
    if certifications_list is None:
        certifications_list = CERTIFICATIONS_LIST
    
    # Sobrescribir referencia si se proporciona custom_reference
    if custom_reference:
        profile_data = profile_data.copy()
        profile_data['reference'] = custom_reference
    
    # Obtener textos según idioma
    texts = TEXTS[language]
    
    doc = Document()
   
    # Page Setup
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
   
    setup_document_styles(doc)
    add_header(doc, logo_path)
    add_footer(doc, profile_data['reference'], language)
   
    # --- A) TÍTULO SUPERIOR ---
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(profile_data['role'])
    set_font_style(run_title, size=14, bold=True, color_rgb=RGBColor(250, 60, 15))
   
    # Agregar espacio extra antes de la tabla de referencia
    doc.add_paragraph()
    
    # --- B) BLOQUE REFERENCIA ---
    table_ref = doc.add_table(rows=1, cols=2)
    table_ref.autofit = False
    table_ref.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Aplicar bordes negros a la tabla
    set_table_borders(table_ref)
   
    # Reference Label (con fondo gris)
    cell_ref_label = table_ref.cell(0, 0)
    set_cell_background(cell_ref_label, "F0F0F0")  # Fondo gris claro
    p = cell_ref_label.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(texts["reference"])
    set_font_style(run, size=10, bold=True)
   
    # Reference Value
    cell_ref_val = table_ref.cell(0, 1)
    p = cell_ref_val.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(profile_data['reference'])
    set_font_style(run, size=10, bold=False)
   
    # Agregar espacio extra después de la tabla
    doc.add_paragraph()
    doc.add_paragraph() # Spacer adicional
   
    # --- SECCIÓN: RESUMEN PROFESIONAL ---
    p_head = doc.add_paragraph()
    run = p_head.add_run(texts["professional_summary"])
    set_font_style(run, size=11, bold=True, underline=True, color_rgb=RGBColor(0, 0, 41)) # #000029
   
    p_sum = doc.add_paragraph(profile_data['summary'])
   
    doc.add_paragraph() # Spacer
   
    # --- SECCIÓN: EXPERIENCIA PROFESIONAL ---
    p_head = doc.add_paragraph()
    run = p_head.add_run(texts["professional_experience"])
    set_font_style(run, size=11, bold=True, underline=True, color_rgb=RGBColor(0, 0, 41))
   
    for exp in experience_list:
        create_experience_table(doc, exp)
        doc.add_paragraph() # Space between tables
       
    # --- SECCIÓN: EDUCACIÓN ---
    p_head = doc.add_paragraph()
    run = p_head.add_run(texts["education"])
    set_font_style(run, size=11, bold=True, underline=True, color_rgb=RGBColor(0, 0, 41))
   
    for edu in education_list:
        p = doc.add_paragraph(style='List Paragraph')
        p.paragraph_format.left_indent = Cm(0.63)
        run = p.add_run(f"• {edu}")
        set_font_style(run, size=11)
       
    doc.add_paragraph()
   
    # --- SECCIÓN: CERTIFICACIONES ---
    p_head = doc.add_paragraph()
    run = p_head.add_run(texts["certifications"])
    set_font_style(run, size=11, bold=True, underline=True, color_rgb=RGBColor(0, 0, 41))
   
    for cert in certifications_list:
        p = doc.add_paragraph(style='List Paragraph')
        p.paragraph_format.left_indent = Cm(0.63)
        run = p.add_run(f"• {cert}")
        set_font_style(run, size=11)

    doc.add_paragraph()
   
    # --- SECCIÓN: IDIOMAS ---
    p_head = doc.add_paragraph()
    run = p_head.add_run(texts["languages"])
    set_font_style(run, size=11, bold=True, underline=True, color_rgb=RGBColor(0, 0, 41))
   
    for lang in profile_data['languages']:
        p = doc.add_paragraph(style='List Paragraph')
        p.paragraph_format.left_indent = Cm(0.63)
        run_bullet = p.add_run("• ")
        set_font_style(run_bullet, size=11)
        run_lang = p.add_run(f"{lang['lang']}: ")
        set_font_style(run_lang, size=11, bold=True)
        run_level = p.add_run(lang['level'])
        set_font_style(run_level, size=11)
       
    doc.add_paragraph()
   
    # --- SECCIÓN: SKILLS ---
    p_head = doc.add_paragraph()
    run = p_head.add_run(texts["skills"])
    set_font_style(run, size=11, bold=True, underline=True, color_rgb=RGBColor(0, 0, 41))
   
    for skill in profile_data['skills']:
        p = doc.add_paragraph(style='List Paragraph')
        p.paragraph_format.left_indent = Cm(0.63)
        run = p.add_run(f"• {skill}")
        set_font_style(run, size=11)

    return doc

# Ejecutar generación
if __name__ == "__main__":
    doc_final = generate_cv('logo.png')
    
    # NOTA: El guardado se hace en main.py usando rutas correctas de Application Support
    # No guardamos aquí para evitar problemas de permisos en macOS
    
    # También guardarlo en BytesIO para uso programático si es necesario
    output_stream = io.BytesIO()
    doc_final.save(output_stream)
    output_stream.seek(0)
    
    print(f"✅ CV generado exitosamente (documento en memoria)")
